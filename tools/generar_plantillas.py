"""
Generador de plantillas Excel del Modelo S-LCA 2.1
==================================================

Crea, en `data/templates/`, los 5 archivos de carga con el MISMO esquema que
lee `src/utils/data_loader.py` (incluida la fila de título donde aplica). Las
etiquetas estructurales vienen rellenas; los VALORES quedan en blanco para que
la persona usuaria los complete.

Uso:
    cd LCA_SOCIAL_2.1
    python tools/generar_plantillas.py
"""
from pathlib import Path

from openpyxl import Workbook

RAIZ = Path(__file__).resolve().parent.parent
DEST = RAIZ / "data" / "templates"
DEST.mkdir(parents=True, exist_ok=True)

DIMENSIONES = [
    ("IA", "Aspiraciones y expectativas"),
    ("IAT", "Tecnología y apropiación"),
    ("ICI", "Confianza institucional"),
    ("IVC", "Identidad cultural"),
    ("IME", "Mentalidad emprendedora"),
    ("IIN", "Normas de género"),
    ("IPRA", "Riesgo climático"),
    ("ICS", "Consumo sostenible"),
]
ACTORES = ["Comunidad Local", "Gobierno Local", "Proveedores Técnicos",
           "Actores Económicos", "Actores Externos"]
INDICADORES = [
    ("Pobreza Multidimensional", "Pobreza multidimensional", "%", "DANE"),
    ("Informalidad", "Tasa de informalidad laboral", "%", "DANE-GEIH"),
    ("Educacion Promedio", "Años promedio de escolaridad", "años", "Admvo. Mpal."),
    ("Acceso Internet", "Cobertura hogares con internet", "%", "MinTIC"),
    ("Desempleo Juvenil", "Desempleo < 25 años", "%", "DANE"),
    ("Participacion Ciudadana", "Participación en espacios cívicos", "%", "Admvo. Mpal."),
]


def _guardar(wb, nombre):
    wb.save(DEST / nombre)
    print("  +", nombre)


def panel_0():
    wb = Workbook(); ws = wb.active; ws.title = "PROTAGONISTAS_READINESS"
    ws.append(["Actor", "Readiness_Actual", "Readiness_Requerido", "Brecha",
               "Historial_12m", "Cambio_12m", "Tendencia", "Volatilidad",
               "Alertas", "Descripcion_Alerta"])
    for actor in ACTORES:
        ws.append([actor, None, None, None, None, None, "", "", "", ""])
    _guardar(wb, "PANEL_0_PROTAGONISTAS.xlsx")


def dimensiones():
    wb = Workbook(); ws = wb.active; ws.title = "RESUMEN"
    ws.append(["DIMENSIONES COGNITIVAS — RESUMEN POR ACTOR Y FASE"])
    ws.append(["Dimensión", "Nombre Completo", "Fase 1", "Fase 2", "Fase 3", "Estado F1"])
    for cod, nombre in DIMENSIONES:
        ws.append([cod, nombre, None, None, None, ""])
    _guardar(wb, "DIMENSIONES_COGNITIVAS.xlsx")


def contexto():
    wb = Workbook(); ws = wb.active; ws.title = "INDICADORES_DANE"
    ws.append(["Indicador", "Nombre", "Valor", "Unidad", "Fuente", "Año", "Comp_Nacional"])
    for cod, nombre, unidad, fuente in INDICADORES:
        ws.append([cod, nombre, None, unidad, fuente, None, None])
    _guardar(wb, "CONTEXTO_TERRITORIAL.xlsx")


def memoria():
    wb = Workbook(); ws = wb.active; ws.title = "SERIE_TEMPORAL_12M"
    ws.append(["Fecha", "IA", "IAT", "ICI", "IVC", "IME", "IIN", "IPRA", "ICS", "Eventos"])
    for i in range(12):
        ws.append([f"Mes {i + 1}", None, None, None, None, None, None, None, None, ""])
    _guardar(wb, "MEMORIA_TEMPORAL.xlsx")


def operativa():
    wb = Workbook(); ws = wb.active; ws.title = "INFRAESTRUCTURA"
    ws.append(["DIMENSIÓN OPERATIVA-MATERIAL"])
    ws.append(["Tipo", "Recurso Específico", "Cantidad", "Estado",
               "Edad (años)", "Capacidad", "Costo Reemplazo (USD)"])
    for tipo in ["Electricidad", "Agua", "Telecomunicaciones", "Vías",
                 "Energía solar", "Almacenamiento", "Sala comunitaria"]:
        ws.append([tipo, "", None, "", None, "", None])
    ws.append([])
    ws.append(["IOM — Índice Operativo (0-100)", "", "", "IOM Calculado",
               None, "Estado:", ""])
    _guardar(wb, "OPERATIVA_MATERIAL.xlsx")


if __name__ == "__main__":
    print("Generando plantillas en", DEST)
    panel_0(); dimensiones(); contexto(); memoria(); operativa()
    print("Listo.")
